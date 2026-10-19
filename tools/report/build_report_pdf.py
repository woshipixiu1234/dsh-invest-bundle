#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""build_report_pdf.py — Markdown 研报 → 中文 A4 DOCX → PDF。

契约见 ``docs/report-pipeline.md`` §1（视觉规范）、§4（Markdown 子集）、
§5（meta schema）、§6（本脚本 CLI 与输出）、§8（质量门槛）。

    python build_report_pdf.py --md report.md --out-dir <dir> --meta report-meta.json \
           [--engine auto|libreoffice|soffice|word] [--title ...] [--name report] \
           [--no-docx|--keep-docx] [--json]

设计约束（刻意遵守）：
* 只依赖 ``python-docx`` + 标准库；同目录 ``palette.py`` 是纯常量模块（可选导入，
  缺失时回退到内置同一套色值），保证 Word 侧与图表侧配色一致。
* 不硬编码任何公司名/日期/数字：品牌话术全部来自 ``--meta``，正文全部来自 ``--md``。
* 版式原语完全参数化（字号/间距/宽度都由常量表驱动），换一份报告只换输入。

转换引擎（``--engine``）：
  ``auto``       依次探测 LibreOffice Kit → 系统 soffice → Word COM，全不可用退出码 1
  ``libreoffice``DSH 自带 LibreOffice Kit（node + cli.js，可用 --lo-node/--lo-cli 或
                 环境变量 DSH_LO_NODE / DSH_LO_CLI 指定）
  ``soffice``    系统/便携版 soffice
  ``word``       Windows Word COM（ExportAsFixedFormat，需要 pywin32）
  ``chrome``     不支持：Chrome 无头打印无法生成 docx 表头/页码域，明确报错而不是静默降级
"""
from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

__version__ = "1.0.0"

# =====================================================================================
# 0. 共享视觉常量（palette.py 优先，缺失则用同一套内置值）
# =====================================================================================
_HERE = os.path.dirname(os.path.abspath(__file__))
if _HERE not in sys.path:
    sys.path.insert(0, _HERE)

try:  # pragma: no cover - 正常运行都会成功
    import palette as _palette

    _DOCX = dict(_palette.DOCX)
    EA_H = getattr(_palette, "EA", "微软雅黑")
    EA_BODY = getattr(_palette, "EA_BODY", "宋体")
    LATIN = getattr(_palette, "LATIN", "Times New Roman")
    LATIN_H = getattr(_palette, "LATIN_H", "微软雅黑")
    _PAGE = dict(getattr(_palette, "PAGE", {}) or {})
except Exception:  # 独立拷贝场景：不因为缺 palette.py 就崩
    _DOCX = {
        "NAVY": "142B47", "BLUE": "1F4E79", "DARK": "222830", "GREY": "5A6470",
        "RED": "A82F22", "AMBER": "9A631E", "BORDER": "C9D2DC", "LIGHT": "F5F8FB",
        "WHITE": "FFFFFF",
    }
    EA_H, EA_BODY, LATIN, LATIN_H = "微软雅黑", "宋体", "Times New Roman", "微软雅黑"
    _PAGE = {}

NAVY = _DOCX.get("NAVY", "142B47")
BLUE = _DOCX.get("BLUE", "1F4E79")
DARK = _DOCX.get("DARK", "222830")
GREY = _DOCX.get("GREY", "5A6470")
RED = _DOCX.get("RED", "A82F22")
AMBER = _DOCX.get("AMBER", "9A631E")
BORDER = _DOCX.get("BORDER", "C9D2DC")
LIGHT = _DOCX.get("LIGHT", "F5F8FB")
WHITE = _DOCX.get("WHITE", "FFFFFF")

# 页面尺寸（cm）——契约 §1
PAGE_W_CM = float(_PAGE.get("width_cm", 21.0))
PAGE_H_CM = float(_PAGE.get("height_cm", 29.7))
MARGIN_T_CM = float(_PAGE.get("margin_top_cm", 2.0))
MARGIN_B_CM = float(_PAGE.get("margin_bottom_cm", 1.8))
MARGIN_L_CM = float(_PAGE.get("margin_left_cm", 2.2))
MARGIN_R_CM = float(_PAGE.get("margin_right_cm", 2.2))
CONTENT_W_CM = PAGE_W_CM - MARGIN_L_CM - MARGIN_R_CM          # 16.6cm
CONTENT_H_CM = PAGE_H_CM - MARGIN_T_CM - MARGIN_B_CM          # 25.9cm

IMG_MAX_W_CM = 15.6            # 契约 §4：插图最大 15.6cm（< 版心 16.6cm，绝不溢出）
IMG_MAX_H_CM = 19.0            # 单图最高，避免一张图把整页顶掉
COVER_IMG_MAX_H_CM = 7.2       # 封面图最高：铺满版心 16.6cm 后超出的部分上下居中裁切（保证封面单页排完）

TWIPS_PER_CM = 566.929         # 1cm = 566.929 twips（dxa）
FONT_MONO = "Consolas"         # 行内代码：Windows 必有，避免 Courier 缺失告警


# =====================================================================================
# 1. OOXML 底层工具
# =====================================================================================
# OOXML 对 pPr/tcPr/tblPr/trPr 的子元素**顺序**有强校验：Word 会直接拒绝乱序文件，
# LibreOffice 宽容但不该依赖。凡是手工 append 过子元素的地方，最后都调用 _fix_order()
# 按 schema 顺序重排，保证同一份 docx 在 LO 引擎与 Word 引擎下都能打开。
_ORDER = {
    "pPr": ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
            "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
            "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
            "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
            "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
            "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
            "sectPr", "pPrChange"],
    "tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap",
             "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark", "headers",
             "cellIns", "cellDel", "cellMerge", "tcPrChange"],
    "tblPr": ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
              "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders",
              "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription",
              "tblPrChange"],
    "trPr": ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter", "cantSplit",
             "trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden", "ins", "del",
             "trPrChange"],
}


def _fix_order(parent, kind):
    """把 parent 的子元素按 schema 顺序重排（append 已存在的元素在 lxml 里等于移动）。"""
    order = _ORDER[kind]
    idx = {qn("w:" + name): i for i, name in enumerate(order)}
    kids = list(parent)
    kids.sort(key=lambda e: idx.get(e.tag, 999))
    for el in kids:
        parent.append(el)
    return parent


def _sub(parent, tag, **attrs):
    """取得或创建子元素，并写入属性。"""
    el = parent.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        parent.append(el)
    for k, v in attrs.items():
        el.set(qn("w:" + k), str(v))
    return el


def rgb(hexstr):
    """'1F4E79' / '#1F4E79' / RGBColor → RGBColor；已是 RGBColor 时原样返回（幂等）。"""
    if isinstance(hexstr, RGBColor):
        return hexstr
    s = str(hexstr or "").strip().lstrip("#")
    if len(s) == 6:
        s = s.upper()
    return RGBColor.from_string(s)


def set_run_font(run, size=10.5, bold=False, color=DARK, ea=None, latin=None,
                 italic=False, mono=False, underline=False):
    """显式写死四组字体槽位，避免落到主题字体（LO 会报 Courier 缺失）。"""
    ea = FONT_MONO if mono else (ea or EA_BODY)
    latin = FONT_MONO if mono else (latin or LATIN)
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = rgb(color) if color else None
    if underline:
        f.underline = True
    else:
        f.underline = False
    f.name = latin
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rf.attrib.pop(qn("w:" + attr), None)
    rf.set(qn("w:ascii"), latin)
    rf.set(qn("w:hAnsi"), latin)
    rf.set(qn("w:eastAsia"), ea)
    rf.set(qn("w:cs"), latin)
    return run


def shade_paragraph(p, fill):
    ppr = p._element.get_or_add_pPr()
    shd = _sub(ppr, "w:shd", val="clear", color="auto", fill=fill)
    _fix_order(ppr, "pPr")
    return shd


def shade_cell(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    _sub(tcpr, "w:shd", val="clear", color="auto", fill=fill)
    _fix_order(tcpr, "tcPr")


def paragraph_borders(p, spec):
    """spec: {"left": (sz_eighths, color, space_pt), "bottom": (...)}

    sz 是 1/8 磅：3pt ≈ sz 24，4mm ≈ 11.34pt ≈ sz 90。
    """
    ppr = p._element.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    for side in ("top", "left", "bottom", "right"):
        if side not in spec:
            continue
        sz, color, space = spec[side]
        el = OxmlElement("w:" + side)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(int(sz)))
        el.set(qn("w:space"), str(int(space)))
        el.set(qn("w:color"), color)
        pbdr.append(el)
    # w:pBdr 内部也要求 top/left/bottom/right 顺序
    order = {"w:top": 0, "w:left": 1, "w:bottom": 2, "w:right": 3}
    for el in sorted(list(pbdr), key=lambda e: order.get(e.tag, 9)):
        pbdr.append(el)
    _fix_order(ppr, "pPr")


def set_table_borders(table, color=BORDER, sz=4, sides=None):
    """默认四边+内部横竖线都用 BORDER 细线；sides 可只保留部分边。"""
    tblpr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        if sides is not None and edge not in sides:
            el.set(qn("w:val"), "none")
            el.set(qn("w:sz"), "0")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "auto")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(int(sz)))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color)
        borders.append(el)
    tblpr.append(borders)
    _fix_order(tblpr, "tblPr")


def set_cell_borders(cell, spec):
    """单元格级边框（结论框的「只有左蓝线」就靠它）。spec: {"left": (sz, color)}。"""
    tcpr = cell._tc.get_or_add_tcPr()
    tb = tcpr.find(qn("w:tcBorders"))
    if tb is None:
        tb = OxmlElement("w:tcBorders")
        tcpr.append(tb)
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:" + edge)
        if edge in spec:
            sz, color = spec[edge]
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(int(sz)))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color)
        else:
            el.set(qn("w:val"), "nil")
        tb.append(el)
    _fix_order(tcpr, "tcPr")


def table_cell_margins(table, top=45, left=90, bottom=45, right=90):
    tblpr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for key, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement("w:" + key)
        el.set(qn("w:w"), str(int(val)))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblpr.append(mar)
    _fix_order(tblpr, "tblPr")


def set_table_layout_fixed(table, width_cm=CONTENT_W_CM):
    tblpr = table._tbl.tblPr
    _sub(tblpr, "w:tblW", w=int(width_cm * TWIPS_PER_CM), type="dxa")
    layout = tblpr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblpr.append(layout)
    layout.set(qn("w:type"), "fixed")
    _fix_order(tblpr, "tblPr")


def mark_header_row(row):
    """表头行：跨页重复（w:tblHeader）+ 行内不断页（w:cantSplit）。"""
    trpr = row._tr.get_or_add_trPr()
    if trpr.find(qn("w:tblHeader")) is None:
        trpr.append(OxmlElement("w:tblHeader"))
    if trpr.find(qn("w:cantSplit")) is None:
        trpr.append(OxmlElement("w:cantSplit"))
    _fix_order(trpr, "trPr")


def mark_no_split(row):
    trpr = row._tr.get_or_add_trPr()
    if trpr.find(qn("w:cantSplit")) is None:
        trpr.append(OxmlElement("w:cantSplit"))
    _fix_order(trpr, "trPr")


def set_paragraph_outline(p, level):
    """给标题写大纲级别，LibreOffice/Word 导出 PDF 时会生成书签。"""
    ppr = p._element.get_or_add_pPr()
    _sub(ppr, "w:outlineLvl", val=str(max(0, min(8, level))))
    _fix_order(ppr, "pPr")


def add_page_field(paragraph, instr=" PAGE ", placeholder="1", size=8.5, bold=True,
                   color=None, ea=None, latin=None):
    """插入 Word 域（PAGE），带 separate + 缓存值，LibreOffice 转换后能算出真页码。"""
    color = BLUE if color is None else color

    def fld(kind):
        run = paragraph.add_run()
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        run._r.append(el)
        set_run_font(run, size, bold, color, ea or EA_H, latin or LATIN_H)
        return run

    fld("begin")
    run = paragraph.add_run()
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    run._r.append(instr_el)
    set_run_font(run, size, bold, color, ea or EA_H, latin or LATIN_H)
    fld("separate")
    set_run_font(paragraph.add_run(placeholder), size, bold, color, ea or EA_H, latin or LATIN_H)
    fld("end")


# =====================================================================================
# 2. 文本清洗 / 富文本
# =====================================================================================
# 表情符号在中文字体里常是豆腐块，统一降级为 CJK 安全符号或直接丢弃。
_EMOJI_MAP = {
    "✅": "◆", "✔": "◆", "☑": "◆", "✓": "◆", "🟢": "◆",
    "🟡": "◇", "🔴": "◇", "⚪": "◇",
    "⚠": "▲", "⚠️": "▲", "🚩": "▲", "📈": "▲", "📉": "▼", "🔥": "▲",
    "⭐": "★", "🌟": "★", "💡": "※", "📌": "※", "🔑": "※", "📎": "※",
    "➡": "→", "➜": "→", "⟶": "→", "→": "→",
}
_EMOJI_DROP = re.compile("[\ufe0f\u200d\u20e3]")


def sanitize(text):
    """返回 (清洗后文本, 被丢弃的 emoji 个数)。"""
    if not text:
        return "", 0
    for k, v in _EMOJI_MAP.items():
        text = text.replace(k, v)
    text = _EMOJI_DROP.sub("", text)
    dropped = 0

    def _drop(m):
        nonlocal dropped
        dropped += 1
        return ""

    # 剩余星平面 emoji（U+1F000 以上）与杂项符号区，中文字体多无字形
    text = re.sub("[\U0001F000-\U0001FAFF\U00002600-\U000026FF\U0001F1E6-\U0001F1FF]", _drop, text)
    return text, dropped


# 行内标记：**粗体** / `代码` / [文字](url) / *斜体*（按此优先级拆分）
_TOKEN_RE = re.compile(
    r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)\s]+\)|\*(?!\*)(?:[^*\n]+?)\*)"
)


def rich(paragraph, text, size=10.5, color=DARK, bold=False, ea=EA_BODY, latin=LATIN):
    """把一段带行内标记的文字写进 paragraph。"""
    for tok in _TOKEN_RE.split(text or ""):
        if not tok:
            continue
        if len(tok) > 4 and tok.startswith("**") and tok.endswith("**"):
            set_run_font(paragraph.add_run(tok[2:-2]), size, True, color, ea, latin)
        elif len(tok) > 2 and tok.startswith("`") and tok.endswith("`"):
            # 行内代码：等宽 + 深蓝，字号略小
            set_run_font(paragraph.add_run(tok[1:-1]), size - 0.5, bold, BLUE, ea, latin,
                         mono=True)
        elif tok.startswith("[") and "](" in tok:
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", tok)
            if m:
                # 契约：链接只做「蓝色下划线文字」，不加超链接关系，避免依赖项
                set_run_font(paragraph.add_run(m.group(1)), size, bold, BLUE, ea, latin,
                             underline=True)
            else:
                set_run_font(paragraph.add_run(tok), size, bold, color, ea, latin)
        elif len(tok) > 2 and tok.startswith("*") and tok.endswith("*") and not tok.startswith("**"):
            set_run_font(paragraph.add_run(tok[1:-1]), size, bold, color, ea, latin,
                         italic=True)
        else:
            set_run_font(paragraph.add_run(tok), size, bold, color, ea, latin)
    return paragraph


def text_units(s):
    """粗略视觉宽度（CJK=2，ASCII=1），用于表格列宽分配。"""
    s = re.sub(r"\*\*|`|\[|\]\([^)]*\)|[*]", "", s or "")
    total = 0.0
    for ch in s:
        if ch == " ":
            total += 0.55
        elif ord(ch) > 0x2000:
            total += 2.0
        else:
            total += 1.0
    return total


_NUM_HEAD = re.compile(r"^[-−+]?\d[\d,，.]*")


def is_numberish(s):
    """判断单元格是否为「数字型」（数字开头 + 短单位），用于列对齐推断。"""
    t = re.sub(r"\*\*|`|\[([^\]]*)\]\([^)]*\)", r"\1", (s or "").strip())
    if not t:
        return False
    m = _NUM_HEAD.match(t)
    if not m:
        return False
    rest = t[m.end():].strip()
    if len(rest) > 6:
        return False
    return not re.search(r"[。；：，、]", rest)


# =====================================================================================
# 3. Markdown 解析（契约 §4 的子集）
# =====================================================================================
class Block:
    __slots__ = ("kind", "text", "items", "rows", "aligns", "level")

    def __init__(self, kind, text="", items=None, rows=None, aligns=None, level=0):
        self.kind = kind
        self.text = text
        self.items = items
        self.rows = rows
        self.aligns = aligns
        self.level = level


_HR_RE = re.compile(r"^\s{0,3}(?:-{3,}|\*{3,}|_{3,})\s*$")
_H_RE = re.compile(r"^(#{1,4})\s+(.*?)\s*#*\s*$")
_TABLE_SEP_RE = re.compile(r"^\s*\|?[\s:|-]*-[\s:|-]*\|?\s*$")
_IMG_RE = re.compile(r"^\s*!\[(.*?)\]\(\s*([^)]+?)\s*\)\s*$")
_UL_RE = re.compile(r"^(\s*)[-*+]\s+(.*)$")
_OL_RE = re.compile(r"^(\s*)(\d+)[.)、]\s+(.*)$")
_FENCE_RE = re.compile(r"^\s*(```|~~~)")


def _split_row(line):
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def _parse_aligns(sep_line):
    aligns = []
    for cell in _split_row(sep_line):
        c = cell.strip()
        left = c.startswith(":")
        right = c.endswith(":")
        if left and right:
            aligns.append("center")
        elif right:
            aligns.append("right")
        elif left:
            aligns.append("left")
        else:
            aligns.append(None)     # None = 自动推断
    return aligns


def parse_markdown(text):
    """把 Markdown 文本切成块列表。返回 (blocks, warnings)。"""
    warnings = []
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # 先摘出强制分页标记，再整体清掉其余 HTML 注释
    text = re.sub(r"<!--\s*pagebreak\s*-->", "\n\x00PAGEBREAK\x00\n", text, flags=re.I)
    text = re.sub(r"<!--.*?-->", "", text, flags=re.S)

    lines = text.split("\n")
    blocks = []
    i, n = 0, len(lines)

    while i < n:
        raw = lines[i]
        s = raw.strip()

        if s == "\x00PAGEBREAK\x00":
            blocks.append(Block("pagebreak"))
            i += 1
            continue
        if not s:
            i += 1
            continue

        # ---- 围栏代码块（契约未列，但研报里常贴配置/公式，顺手支持）----
        if _FENCE_RE.match(raw):
            fence = _FENCE_RE.match(raw).group(1)
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith(fence):
                buf.append(lines[i])
                i += 1
            i += 1
            blocks.append(Block("code", items=buf))
            continue

        # ---- 标题 ----
        m = _H_RE.match(s)
        if m:
            lvl = len(m.group(1))
            blocks.append(Block("h%d" % lvl, text=m.group(2).strip()))
            i += 1
            continue

        # ---- 分隔线 ----
        if _HR_RE.match(s):
            blocks.append(Block("hr"))
            i += 1
            continue

        # ---- 表格：当前行以 | 开头且下一行是分隔行 ----
        if s.startswith("|") and i + 1 < n and "|" in lines[i + 1] \
                and _TABLE_SEP_RE.match(lines[i + 1]):
            rows, aligns = [_split_row(lines[i])], _parse_aligns(lines[i + 1])
            i += 2
            while i < n and lines[i].strip().startswith("|"):
                rows.append(_split_row(lines[i]))
                i += 1
            widths = {len(r) for r in rows}
            if len(widths) > 1:
                warnings.append("表格列数不一致（%s），已按首行 %d 列对齐"
                                % (sorted(widths), len(rows[0]) if rows else 0))
            blocks.append(Block("table", rows=rows, aligns=aligns))
            continue

        # ---- 插图 ----
        m = _IMG_RE.match(s)
        if m:
            blocks.append(Block("image", text=m.group(2), items=[m.group(1).strip()]))
            i += 1
            continue

        # ---- 结论框 / 引用 ----
        if s.startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            # 引用块内的空行在视觉上就是段落分隔，保留为空串
            blocks.append(Block("quote", items=buf))
            continue

        # ---- 无序列表 ----
        m = _UL_RE.match(raw)
        if m:
            items = []
            while i < n:
                mm = _UL_RE.match(lines[i])
                if not mm:
                    break
                indent = len(mm.group(1).replace("\t", "    "))
                items.append((1 if indent >= 2 else 0, mm.group(2).strip()))
                i += 1
            blocks.append(Block("ul", items=items))
            continue

        # ---- 有序列表 ----
        m = _OL_RE.match(raw)
        if m:
            items = []
            while i < n:
                mm = _OL_RE.match(lines[i])
                if not mm:
                    break
                indent = len(mm.group(1).replace("\t", "    "))
                items.append((1 if indent >= 2 else 0, mm.group(2).strip(), mm.group(3).strip()))
                i += 1
            blocks.append(Block("ol", items=items))
            continue

        # ---- 普通段落：连续纯文本行合并（中文换行不加空格）----
        buf = [s]
        i += 1
        while i < n:
            nxt = lines[i]
            ns = nxt.strip()
            if (not ns or _H_RE.match(ns) or _HR_RE.match(ns) or ns.startswith("|")
                    or ns.startswith(">") or _UL_RE.match(nxt) or _OL_RE.match(nxt)
                    or _IMG_RE.match(ns) or _FENCE_RE.match(nxt)
                    or ns == "\x00PAGEBREAK\x00"):
                break
            buf.append(ns)
            i += 1
        merged = buf[0]
        for part in buf[1:]:
            gap = "" if (merged and ord(merged[-1]) > 0x2000 and part
                         and ord(part[0]) > 0x2000) else " "
            merged += gap + part
        blocks.append(Block("p", text=merged))

    return blocks, warnings


# =====================================================================================
# 4. 文档渲染
# =====================================================================================
class Renderer:
    """把块列表渲染进 python-docx 文档。"""

    def __init__(self, doc, meta, base_dir, warnings, img_max_w_cm=IMG_MAX_W_CM,
                 h2_pagebreak=False):
        self.doc = doc
        self.meta = meta or {}
        self.base_dir = base_dir
        self.warnings = warnings
        self.img_max_w_cm = img_max_w_cm
        self.h2_pagebreak = bool(h2_pagebreak)  # 契约 §4 的「## 段前分页」开关，默认关
        self.at_page_start = False   # 当前是否刚好在页首（避免封面后双重分页）
        self.pending_break = False   # 下一个块之前必须分页
        self.figures = 0

    # ---------------- 低层 ----------------
    def _par(self, space_before=0, space_after=5, line=1.45,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY, keep_next=False, keep_together=False):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(space_before)
        pf.space_after = Pt(space_after)
        pf.line_spacing = line
        pf.keep_with_next = keep_next
        pf.keep_together = keep_together
        if align is not None:
            p.alignment = align
        if self.pending_break:
            pf.page_break_before = True
            self.pending_break = False
            self.at_page_start = True
        return p

    def _finish(self):
        """每个内容块渲染完成后调用：这一页已经不是「页首」了。"""
        self.at_page_start = False

    def _spacer(self, pt=6):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        set_run_font(p.add_run(""), max(1.0, pt))

    def _ensure_pending_before_object(self):
        """表格/图片这类「对象」不能自带 pageBreakBefore，用空段实现分页。"""
        if not self.pending_break:
            return
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.page_break_before = True
        set_run_font(p.add_run(""), 1)
        self.pending_break = False
        self.at_page_start = True

    # ---------------- 标题 ----------------
    def h1(self, text):
        """正文 H1 一般不出现（契约 §4）；出现时降级为章节前的大标题。"""
        p = self._par(space_before=6, space_after=8, line=1.25,
                      align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=True, keep_together=True)
        set_run_font(p.add_run(text), 16, True, NAVY, EA_H, LATIN_H)
        set_paragraph_outline(p, 0)
        self._finish()

    def h2(self, text):
        """H2：15pt NAVY + 左侧 4mm 竖线 + 下边框金线。

        契约 §4 写的是「## 段前分页」，但无条件分页会把章节尾部留成半空页
        （python-docx 量不出剩余高度），所以默认改为「段前 18pt + keep_with_next」
        的连续排版，硬分页交给 `<!-- pagebreak -->`；需要旧行为时传
        h2_pagebreak=True（CLI: --h2-pagebreak always）。
        """
        p = self._par(space_before=18, space_after=9, line=1.2,
                      align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=True, keep_together=True)
        if self.h2_pagebreak and not self.at_page_start \
                and not p.paragraph_format.page_break_before:
            p.paragraph_format.page_break_before = True
            self.at_page_start = True
        p.paragraph_format.left_indent = Cm(0.30)
        # 4mm ≈ 11.34pt ≈ sz 90（1/8 磅）；下边框 2pt 金线压住整行
        paragraph_borders(p, {"left": (90, NAVY, 5), "bottom": (16, AMBER, 4)})
        # 编号（中文数字或阿拉伯数字）用 AMBER 强调，其余 NAVY；文本内容不改
        m = re.match(r"^([一二三四五六七八九十百]+[、.]|\d+(?:\.\d+)*[、.]?\s*|[A-Z][、.])\s*(.*)$",
                     text)
        if m and m.group(2):
            set_run_font(p.add_run(m.group(1).strip() + " "), 15, True, AMBER, EA_H, LATIN_H)
            set_run_font(p.add_run(m.group(2)), 15, True, NAVY, EA_H, LATIN_H)
        else:
            set_run_font(p.add_run(text), 15, True, NAVY, EA_H, LATIN_H)
        set_paragraph_outline(p, 1)
        self._finish()

    def h3(self, text):
        p = self._par(space_before=11, space_after=5, line=1.25,
                      align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=True, keep_together=True)
        set_run_font(p.add_run(text), 11.5, True, BLUE, EA_H, LATIN_H)
        set_paragraph_outline(p, 2)
        self._finish()

    def h4(self, text):
        p = self._par(space_before=8, space_after=4, line=1.3,
                      align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=True, keep_together=True)
        set_run_font(p.add_run(text), 10.5, True, DARK, EA_H, LATIN_H)
        set_paragraph_outline(p, 3)
        self._finish()

    # ---------------- 正文 ----------------
    def paragraph(self, text):
        # 「**表 1　xxx**」独立成行时按表标题排（9pt NAVY，与下方表格同页）
        m = re.match(r"^\*\*\s*((?:表|图|图表|附表)\s*\d+[^*]{0,80})\*\*\s*$", text.strip())
        if m:
            p = self._par(space_before=8, space_after=3, line=1.25,
                          align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=True, keep_together=True)
            set_run_font(p.add_run(m.group(1).strip()), 9, True, NAVY, EA_H, LATIN_H)
            self._finish()
            return
        p = self._par(space_before=0, space_after=6, line=1.45)
        rich(p, text, 10.5, DARK)
        # 「以下是三条：」这类引导句跟后文保持同页
        if (text.endswith("：") or text.endswith(":")) and text_units(text) < 60:
            p.paragraph_format.keep_with_next = True
        self._finish()

    def bullets(self, items):
        for level, text in items:
            indent = 0.62 + level * 0.62
            p = self._par(space_before=0, space_after=3, line=1.4,
                          align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            p.paragraph_format.left_indent = Cm(indent + 0.48)
            p.paragraph_format.first_line_indent = Cm(-0.48)
            bullet = "•" if level == 0 else "–"
            set_run_font(p.add_run(bullet + " "), 10.5, level == 0,
                         BLUE if level == 0 else GREY)
            rich(p, text, 10.5, DARK)
        self._finish()

    def ordered(self, items):
        for level, num, text in items:
            indent = 0.62 + level * 0.62
            p = self._par(space_before=0, space_after=3, line=1.4,
                          align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            p.paragraph_format.left_indent = Cm(indent + 0.62)
            p.paragraph_format.first_line_indent = Cm(-0.62)
            set_run_font(p.add_run("%s. " % num), 10.5, True, BLUE)
            rich(p, text, 10.5, DARK)
        self._finish()

    def quote(self, lines):
        """结论框：浅底 + 左侧 3pt BLUE 边线；首行标签自动加粗（契约 §4）。"""
        lines = [ln for ln in lines]
        while lines and not lines[0]:
            lines.pop(0)
        while lines and not lines[-1]:
            lines.pop()
        if not lines:
            return
        self._ensure_pending_before_object()
        box = self.doc.add_table(rows=1, cols=1)
        box.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_layout_fixed(box, CONTENT_W_CM)
        set_table_borders(box, BORDER, 4, sides=())
        table_cell_margins(box, top=110, left=170, bottom=110, right=150)
        cell = box.cell(0, 0)
        shade_cell(cell, LIGHT)
        set_cell_borders(cell, {"left": (24, BLUE)})     # 24/8 = 3pt
        mark_no_split(box.rows[0])

        first_slot = True
        first_content = True
        for ln in lines:
            p = cell.paragraphs[0] if first_slot else cell.add_paragraph()
            first_slot = False
            pf = p.paragraph_format
            pf.space_before = Pt(0)
            pf.space_after = Pt(3)
            pf.line_spacing = 1.4
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if not ln:
                pf.space_after = Pt(1)
                set_run_font(p.add_run(""), 4)
                continue
            txt = ln
            # 首行「核心结论：…」这类标签自动加粗（契约 §4）
            if first_content and "**" not in txt:
                m = re.match(r"^([^：:]{2,14})[：:]\s*(.*)$", txt)
                if m:
                    set_run_font(p.add_run(m.group(1) + "："), 9.8, True, NAVY, EA_H, LATIN_H)
                    txt = m.group(2)
            first_content = False
            rich(p, txt, 9.8, DARK)
        self._spacer(6)
        self._finish()

    def divider(self):
        p = self._par(space_before=6, space_after=8, line=1.0, align=None)
        paragraph_borders(p, {"bottom": (8, BORDER, 1)})
        set_run_font(p.add_run(""), 2)
        self._finish()

    def code_block(self, lines):
        self._ensure_pending_before_object()
        box = self.doc.add_table(rows=1, cols=1)
        box.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_layout_fixed(box, CONTENT_W_CM)
        set_table_borders(box, BORDER, 4)
        table_cell_margins(box, top=90, left=150, bottom=90, right=120)
        cell = box.cell(0, 0)
        shade_cell(cell, "F7F9FB")
        first = True
        for ln in lines:
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.15
            set_run_font(p.add_run(ln if ln else " "), 8.2, False, "24292F", mono=True)
        self._spacer(6)
        self._finish()

    # ---------------- 表格 ----------------
    def table(self, rows, aligns):
        if not rows:
            return
        ncol = max(len(r) for r in rows)
        norm = []
        for r in rows:
            r = list(r)
            if len(r) < ncol:
                r += [""] * (ncol - len(r))
            norm.append(r[:ncol])
        header, body = norm[0], norm[1:]
        fs = self._table_font(ncol)

        # 列对齐：显式分隔行 > 数字列自动右对齐
        col_align = []
        for ci in range(ncol):
            explicit = aligns[ci] if ci < len(aligns) else None
            if explicit:
                col_align.append(explicit)
                continue
            vals = [r[ci] for r in body if r[ci].strip()]
            hits = sum(1 for v in vals if is_numberish(v))
            auto = "right" if vals and hits / len(vals) >= 0.6 else "left"
            # 首列表头是文字（如「项目」「情景」「电价变动＼电量变动」）时，
            # 它就是行标签列：即使格子里全是数字也左对齐，否则会贴到列右边。
            if ci == 0 and not is_numberish(header[ci].replace("**", "")):
                auto = "left"
            col_align.append(auto)

        self._ensure_pending_before_object()
        t = self.doc.add_table(rows=len(norm), cols=ncol)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        set_table_layout_fixed(t, CONTENT_W_CM)
        set_table_borders(t, BORDER, 4)
        table_cell_margins(t, top=42, left=80, bottom=42, right=80)

        widths = self._col_widths(norm, CONTENT_W_CM)
        for ci, w in enumerate(widths):
            t.columns[ci].width = Cm(w)

        for ri, row in enumerate(norm):
            is_header = ri == 0
            for ci in range(ncol):
                cell = t.cell(ri, ci)
                cell.width = Cm(widths[ci])
                if is_header:
                    shade_cell(cell, NAVY)
                elif ri % 2 == 0:                      # 偶数行斑马（契约 §4）
                    shade_cell(cell, LIGHT)
                p = cell.paragraphs[0]
                pf = p.paragraph_format
                pf.space_before = Pt(1)
                pf.space_after = Pt(1)
                pf.line_spacing = 1.15
                if is_header:
                    # 表头对齐跟随列：文字列左对齐、数字列居中（卖方报告惯例）
                    p.alignment = (WD_ALIGN_PARAGRAPH.LEFT if col_align[ci] == "left"
                                   else WD_ALIGN_PARAGRAPH.CENTER)
                    rich(p, row[ci].replace("**", ""), fs, rgb(WHITE), True, EA_H, LATIN_H)
                else:
                    p.alignment = {"right": WD_ALIGN_PARAGRAPH.RIGHT,
                                   "center": WD_ALIGN_PARAGRAPH.CENTER,
                                   "left": WD_ALIGN_PARAGRAPH.LEFT}[col_align[ci]]
                    rich(p, row[ci], fs, DARK)
            if is_header:
                mark_header_row(t.rows[ri])           # 长表跨页重复表头
            else:
                mark_no_split(t.rows[ri])
        self._spacer(7)
        self._finish()

    @staticmethod
    def _table_font(ncol):
        if ncol <= 5:
            return 8.5
        if ncol <= 7:
            return 7.8
        return 7.2

    @staticmethod
    def _col_widths(rows, total_cm):
        ncol = len(rows[0])
        weights = []
        for ci in range(ncol):
            head = text_units(rows[0][ci])
            cells = sorted(text_units(r[ci]) for r in rows[1:]) or [0.0]
            p70 = cells[int(0.7 * (len(cells) - 1))]
            p90 = cells[int(0.9 * (len(cells) - 1))]
            # 表头只做「轻微加权」：正文很短的列（如敏感性表的行标签列）不允许
            # 因为表头写着「电价变动＼电量变动」就吃掉半张表，让表头自然折行。
            head = min(head, max(p90 * 1.9, 6.0))
            weights.append(min(42.0, max(head * 1.05, p70 * 0.78, 3.2)))
        total_w = sum(weights) or 1.0
        widths = [total_cm * w / total_w for w in weights]
        floor = 1.45 if ncol <= 6 else 1.15
        for _ in range(3):                            # 抬到最小宽度后重新分配剩余宽度
            deficit = sum(floor - w for w in widths if w < floor)
            if deficit <= 0:
                break
            spare = [i for i, w in enumerate(widths) if w > floor]
            pool = sum(widths[i] - floor for i in spare) or 1.0
            for i in range(len(widths)):
                widths[i] = floor if widths[i] < floor else widths[i]
            for i in spare:
                widths[i] -= deficit * (widths[i] - floor) / pool
        return widths

    # ---------------- 插图 ----------------
    def image(self, path, caption):
        real = self._resolve_asset(path)
        if not real or not os.path.isfile(real):
            self.warnings.append("图片缺失，已用占位文字代替：%s" % path)
            p = self._par(space_before=6, space_after=6, line=1.2,
                          align=WD_ALIGN_PARAGRAPH.CENTER)
            set_run_font(p.add_run("［缺失图片］%s" % path), 9, False, RED)
            self._finish()
            return
        self._ensure_pending_before_object()
        w_cm, h_cm = self._image_size(real)
        p = self._par(space_before=8, space_after=3, line=1.0,
                      align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
        try:
            p.add_run().add_picture(real, width=Cm(w_cm), height=Cm(h_cm))
            self.figures += 1
        except Exception as exc:                       # 坏图不该毁掉整份报告
            self.warnings.append("图片插入失败（%s）：%s" % (path, exc))
            set_run_font(p.add_run("［图片无法插入］%s" % path), 9, False, RED)
        if caption:
            cap = self._par(space_before=0, space_after=11, line=1.25,
                            align=WD_ALIGN_PARAGRAPH.CENTER)
            set_run_font(cap.add_run(caption), 8.5, False, GREY, EA_H, LATIN_H)
        self._finish()

    def _resolve_asset(self, rel):
        rel = (rel or "").strip().replace("\\", "/")
        if not rel:
            return None
        if os.path.isabs(rel):
            return rel if os.path.isfile(rel) else None
        bases = [self.base_dir, self.meta.get("_meta_dir") or "", os.getcwd()]
        for base in bases:
            if not base:
                continue
            cand = os.path.normpath(os.path.join(base, rel))
            if os.path.isfile(cand):
                return cand
        return None

    def _image_size(self, path):
        """按像素/DPI 还原原始尺寸，再限制在版心内且等比缩放。"""
        px_w = px_h = None
        dpi = 96.0
        try:
            from PIL import Image                    # 有就用，精度更高
            with Image.open(path) as im:
                px_w, px_h = im.size
                dpi_info = im.info.get("dpi")
                if dpi_info and dpi_info[0]:
                    dpi = float(dpi_info[0])
        except Exception:
            if px_w is None:
                try:
                    from docx.image.image import Image as DxImage
                    img = DxImage.from_file(path)
                    px_w, px_h = img.px_width, img.px_height
                    dpi = float(img.horz_dpi or 96.0)
                except Exception:
                    return self.img_max_w_cm, self.img_max_w_cm * 0.6
        if not px_w or not px_h:
            return self.img_max_w_cm, self.img_max_w_cm * 0.6
        cm_w = px_w / dpi * 2.54
        cm_h = px_h / dpi * 2.54
        scale = 1.0
        if cm_w > self.img_max_w_cm:
            scale = min(scale, self.img_max_w_cm / cm_w)
        if cm_h * scale > IMG_MAX_H_CM:
            scale = min(scale, IMG_MAX_H_CM / cm_h)
        if scale >= 1.0:
            # 不放大：小图按原始尺寸居中，避免糊
            return cm_w, cm_h
        return cm_w * scale, cm_h * scale



# =====================================================================================
# 5. 封面 / 页脚 / 文档装配
# =====================================================================================
def _purge_missing_fonts(doc):
    """python-docx 默认模板里带一个用 Courier 的字符样式（HTML Code 一类），
    Windows 只有 Courier New 没有 Courier，LibreOffice 转换时会报「缺失字体」。
    这里把样式表里所有 Courier 引用改写成 FONT_MONO，并删掉 fontTable 里的声明。
    """
    fixed, dropped = 0, 0
    root = doc.styles.element
    for rf in root.iter(qn("w:rFonts")):
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            val = (rf.get(qn("w:" + attr)) or "").strip().lower()
            if val in ("courier", "courier new"):
                rf.set(qn("w:" + attr), FONT_MONO)
                fixed += 1
    try:
        for part in doc.part.package.iter_parts():
            pn = str(part.partname)
            if not (pn.endswith("fontTable.xml") or pn.endswith("stylesWithEffects.xml")):
                continue
            # 1) 能拿到解析后元素就直接删 <w:font w:name="Courier">
            el = getattr(part, "element", None) or getattr(part, "_element", None)
            if el is not None:
                for fnt in list(el.iter(qn("w:font"))):
                    if (fnt.get(qn("w:name")) or "").strip().lower() in ("courier", "courier new"):
                        fnt.getparent().remove(fnt)
                        dropped += 1
            # 2) 纯字节部件（stylesWithEffects）直接做文本替换
            try:
                blob = part.blob
            except Exception:
                blob = b""
            if b"Courier" in blob:
                part._blob = blob.replace(b"Courier New", b"Consolas").replace(b"Courier", b"Consolas")
                dropped += 1
    except Exception:
        pass
    return fixed, dropped


def patch_doc_defaults(doc):
    """把 docDefaults 的主题字体换成显式字体，避免 Word/LO 回落到 Courier/Calibri。"""
    _purge_missing_fonts(doc)
    styles_el = doc.styles.element
    defaults = styles_el.find(qn("w:docDefaults"))
    if defaults is None:
        return
    rpr_default = defaults.find(qn("w:rPrDefault"))
    if rpr_default is None:
        return
    rpr = rpr_default.find(qn("w:rPr"))
    if rpr is None:
        return
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rf.attrib.pop(qn("w:" + attr), None)
    rf.set(qn("w:ascii"), LATIN)
    rf.set(qn("w:hAnsi"), LATIN)
    rf.set(qn("w:eastAsia"), EA_BODY)
    rf.set(qn("w:cs"), LATIN)
    sz = rpr.find(qn("w:sz"))
    if sz is None:
        sz = OxmlElement("w:sz")
        rpr.append(sz)
    sz.set(qn("w:val"), "21")           # 10.5pt = 21 半磅


def setup_document(meta):
    doc = Document()
    patch_doc_defaults(doc)

    sec = doc.sections[0]
    sec.page_width = Cm(PAGE_W_CM)
    sec.page_height = Cm(PAGE_H_CM)
    sec.top_margin = Cm(MARGIN_T_CM)
    sec.bottom_margin = Cm(MARGIN_B_CM)
    sec.left_margin = Cm(MARGIN_L_CM)
    sec.right_margin = Cm(MARGIN_R_CM)
    sec.header_distance = Cm(1.2)
    sec.footer_distance = Cm(0.9)
    # 封面不要页脚（页码仍在全文档连续计数）
    sec.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = LATIN
    normal.font.size = Pt(10.5)
    normal.element.get_or_add_rPr()
    rf = normal.element.rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        normal.element.rPr.insert(0, rf)
    rf.set(qn("w:eastAsia"), EA_BODY)
    pf = normal.paragraph_format
    pf.line_spacing = 1.45
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    return doc, sec


def footer_text(meta):
    custom = (meta.get("footer") or "").strip()
    if custom:
        return custom
    title = (meta.get("title") or "").strip()
    ticker = (meta.get("ticker") or "").strip()
    head = title
    if ticker:
        head = "%s（%s）" % (title, ticker) if title else ticker
    return ("%s · 深度研究 · 不构成投资建议" % head).strip(" ·")


def build_footer(sec, meta):
    """页脚：左侧短标题+免责，右侧页码（PAGE 域），中间细线压顶。"""
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = fp.paragraph_format
    pf.space_before = Pt(2)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    # 默认模板的 Footer 样式自带 4680/9360 两个制表位（居中/右），会让 \t 提前落位，
    # 必须先把样式里的制表位清空，再挂自己的右对齐制表位。
    try:
        st = fp.style
        if st is not None:
            st.paragraph_format.tab_stops.clear_all()
    except Exception:
        pass
    pf.tab_stops.add_tab_stop(Cm(CONTENT_W_CM), WD_TAB_ALIGNMENT.RIGHT)
    paragraph_borders(fp, {"top": (6, BORDER, 6)})

    set_run_font(fp.add_run(footer_text(meta)), 7.5, False, GREY, EA_BODY, LATIN)
    set_run_font(fp.add_run("\t"), 7.5, False, GREY, EA_BODY, LATIN)
    add_page_field(fp, " PAGE ", "1", size=8.5, bold=True, color=BLUE)


def _core_box(doc, core, width_cm, title_default="核心结论", para_size=9.5):
    """浅底 + 左侧 3pt BLUE 边线的结论框（封面与正文共用同一套观感）。"""
    box = doc.add_table(rows=1, cols=1)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_layout_fixed(box, width_cm)
    set_table_borders(box, BORDER, 4, sides=())
    table_cell_margins(box, top=130, left=180, bottom=130, right=160)
    cell = box.cell(0, 0)
    shade_cell(cell, LIGHT)
    set_cell_borders(cell, {"left": (24, BLUE)})
    mark_no_split(box.rows[0])

    title = (core.get("title") or title_default).strip()
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.2
    set_run_font(p.add_run(title), 9.5, True, AMBER, EA_H, LATIN_H)

    for para in (core.get("paragraphs") or [])[:3]:
        text, _ = sanitize(para)
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.42
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        rich(p, text, para_size, DARK)

    verdict = (core.get("verdict") or "").strip()
    if verdict:
        p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_run_font(p.add_run(sanitize(verdict)[0]), 9.5, True, RED, EA_H, LATIN_H)
    return box


def _crop_picture_vertical(shape, crop_frac):
    """按比例上下对称裁切图片（crop_frac 是每侧占原图高度的比例），保持显示框不变。"""
    if crop_frac <= 0:
        return
    inline = shape._inline
    blip_fill = inline.graphic.graphicData.pic.blipFill
    src = blip_fill.find(qn("a:srcRect"))
    if src is None:
        src = OxmlElement("a:srcRect")
        blip = blip_fill.find(qn("a:blip"))
        blip.addnext(src)
    v = str(int(round(crop_frac * 100000)))
    src.set("t", v)
    src.set("b", v)


def build_cover(doc, meta, base_dir, warnings, img_max_w_cm=IMG_MAX_W_CM):
    """封面：有 meta.cover 时放全宽封面图，否则纯排版封面。两种都单页排完。"""
    title = (meta.get("title") or "").strip() or "研究报告"
    ticker = (meta.get("ticker") or "").strip()
    kicker = (meta.get("kicker") or "").strip()
    subtitle = (meta.get("subtitle") or "").strip()
    as_of = (meta.get("as_of") or "").strip()
    sources_asof = (meta.get("sources_asof") or "").strip()
    disclaimer = (meta.get("disclaimer") or "本报告不构成投资建议。投资有风险，入市需谨慎。").strip()

    cover = meta.get("cover") or {}
    cover_path = None
    if isinstance(cover, dict) and cover.get("path"):
        cover_path = cover["path"]
    elif isinstance(cover, str) and cover.strip():
        cover_path = cover.strip()
    cover_real = None
    if cover_path:
        for base in (base_dir, meta.get("_meta_dir") or "", os.getcwd()):
            if not base:
                continue
            cand = os.path.normpath(os.path.join(base, cover_path))
            if os.path.isfile(cand):
                cover_real = cand
                break
        if not cover_real:
            warnings.append("封面图不存在，已回退为排版封面：%s" % cover_path)

    # ---- 封面图 ----
    if cover_real:
        # 封面图永远铺满版心宽度（16.6cm），高度超过上限时上下居中裁切，
        # 而不是整体缩小（缩小会留下大片白边）。不读 DPI 元数据。
        w_cm, h_cm, crop_frac = CONTENT_W_CM, CONTENT_W_CM * 0.45, 0.0
        try:
            from PIL import Image
            with Image.open(cover_real) as im:
                px_w, px_h = im.size
            if px_w > 0 and px_h > 0:
                w_cm = CONTENT_W_CM
                h_cm = w_cm * px_h / px_w
                if h_cm > COVER_IMG_MAX_H_CM:
                    crop_frac = (h_cm - COVER_IMG_MAX_H_CM) / h_cm / 2.0
                    h_cm = COVER_IMG_MAX_H_CM
        except Exception:
            pass
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(12)
        run = p.add_run()
        shape = run.add_picture(cover_real, width=Cm(w_cm), height=Cm(h_cm))
        _crop_picture_vertical(shape, crop_frac)
        top_gap = 0
    else:
        top_gap = 58

    # ---- kicker ----
    if kicker:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(top_gap)
        p.paragraph_format.space_after = Pt(4)
        set_run_font(p.add_run(kicker), 10.5, True, AMBER, EA_H, LATIN_H)
        top_gap = 0
    if top_gap:
        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(top_gap)
        sp.paragraph_format.space_after = Pt(0)
        set_run_font(sp.add_run(""), 2)

    # ---- 大标题 + 代码 ----
    size = 30 if len(title) <= 8 else (26 if len(title) <= 12 else 22)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.12
    set_run_font(p.add_run(title), size, True, NAVY, EA_H, LATIN_H)
    if ticker:
        set_run_font(p.add_run("　" + ticker), max(12.0, size * 0.42), False, GREY, EA_H, LATIN_H)

    # ---- 金线 ----
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.0
    paragraph_borders(p, {"bottom": (18, AMBER, 2)})
    set_run_font(p.add_run(""), 2)

    # ---- 副标题 ----
    if subtitle:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(14)
        p.paragraph_format.line_spacing = 1.3
        set_run_font(p.add_run(subtitle), 11, False, BLUE, EA_H, LATIN_H)

    # ---- 指标表（2–4 项）----
    metrics = [m for m in (meta.get("metrics") or []) if isinstance(m, dict)][:4]
    if len(meta.get("metrics") or []) > 4:
        warnings.append("meta.metrics 超过 4 项，封面只排前 4 项（契约 §5）")
    if metrics:
        ncol = len(metrics)
        t = doc.add_table(rows=2, cols=ncol)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        set_table_layout_fixed(t, CONTENT_W_CM)
        set_table_borders(t, "DDE3EA", 4)
        table_cell_margins(t, top=90, left=60, bottom=90, right=60)
        col_w = CONTENT_W_CM / ncol
        for ci in range(ncol):
            t.columns[ci].width = Cm(col_w)
            c0, c1 = t.cell(0, ci), t.cell(1, ci)
            c0.width = c1.width = Cm(col_w)
            shade_cell(c0, "F1F5F9")
            p0 = c0.paragraphs[0]
            p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p0.paragraph_format.space_before = Pt(1)
            p0.paragraph_format.space_after = Pt(1)
            p0.paragraph_format.line_spacing = 1.1
            set_run_font(p0.add_run(str(metrics[ci].get("label", ""))), 7.5, False, GREY,
                         EA_H, LATIN_H)
            p1 = c1.paragraphs[0]
            p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p1.paragraph_format.space_before = Pt(1)
            p1.paragraph_format.space_after = Pt(1)
            p1.paragraph_format.line_spacing = 1.15
            set_run_font(p1.add_run(str(metrics[ci].get("value", ""))), 12.5, True, NAVY,
                         EA_H, LATIN_H)
        mark_no_split(t.rows[0])
        mark_no_split(t.rows[1])
        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(0)
        set_run_font(sp.add_run(""), 8)

    # ---- 核心结论框 ----
    core = meta.get("core") or {}
    if core.get("paragraphs") or core.get("verdict"):
        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(0)
        set_run_font(sp.add_run(""), 10)
        _core_box(doc, core, CONTENT_W_CM, para_size=9.5)

    # ---- 报告日期 / 数据基准 ----
    # 封面底部留白太多会显得没做完：用一条细分隔线把「报告信息带」压到页面下方，
    # 分隔线以上保持呼吸感。有封面图时空间紧张，呼吸留白缩小。
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(8 if cover_real else 22)
    sp.paragraph_format.space_after = Pt(0)
    set_run_font(sp.add_run(""), 2)
    if as_of or sources_asof:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.3
        paragraph_borders(p, {"top": (6, BORDER, 6)})
        if as_of:
            set_run_font(p.add_run("报告日期 "), 8, False, GREY, EA_BODY, LATIN)
            set_run_font(p.add_run(as_of), 8, True, NAVY, EA_BODY, LATIN)
        if sources_asof:
            set_run_font(p.add_run(("　｜　数据基准：" if as_of else "数据基准：")), 8, False,
                         GREY, EA_BODY, LATIN)
            set_run_font(p.add_run(sources_asof), 8, False, GREY, EA_BODY, LATIN)
    if disclaimer:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.3
        set_run_font(p.add_run(disclaimer), 8, False, RED, EA_BODY, LATIN)


def build_document(meta, blocks, base_dir, warnings, img_max_w_cm=IMG_MAX_W_CM,
                   h2_pagebreak=False):
    # 契约 §4：`# ` 是文档大标题——meta 没给 title 时用它当封面标题，
    # 给了 title 就把它当作封面的重复不再排版，避免封面上出现两遍同名大标题。
    if blocks and blocks[0].kind == "h1":
        if not (meta.get("title") or "").strip():
            meta["title"] = blocks[0].text
        blocks = blocks[1:]

    doc, sec = setup_document(meta)
    build_footer(sec, meta)
    build_cover(doc, meta, base_dir, warnings, img_max_w_cm)

    r = Renderer(doc, meta, base_dir, warnings, img_max_w_cm, h2_pagebreak)
    r.pending_break = True          # 封面之后强制从新页开始正文
    for b in blocks:
        kind = b.kind
        if kind == "h1":
            r.h1(b.text)
        elif kind == "h2":
            r.h2(b.text)
        elif kind == "h3":
            r.h3(b.text)
        elif kind == "h4":
            r.h4(b.text)
        elif kind == "p":
            r.paragraph(b.text)
        elif kind == "ul":
            r.bullets(b.items)
        elif kind == "ol":
            r.ordered(b.items)
        elif kind == "table":
            r.table(b.rows, b.aligns)
        elif kind == "image":
            r.image(b.text, (b.items or [""])[0])
        elif kind == "quote":
            r.quote(b.items)
        elif kind == "hr":
            r.divider()
        elif kind == "code":
            r.code_block(b.items)
        elif kind == "pagebreak":
            r.pending_break = True
    return doc, r.figures


# =====================================================================================
# 6. 转换引擎
# =====================================================================================
LO_NODE_DEFAULT = r"D:\deepseek harness\resources\runtime\primary-runtime\dependencies\node\bin\node.exe"
LO_CLI_DEFAULT = (r"D:\deepseek harness\resources\app.asar.unpacked\dsh\node_modules"
                  r"\@deepseek-ai\libreoffice-kit\lib\cli.js")
SOFFICE_CANDIDATES = (
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    "/usr/bin/soffice", "/usr/local/bin/soffice",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
)


class EngineError(RuntimeError):
    pass


def _probe_lo_kit(node_hint=None, cli_hint=None):
    node = node_hint or os.environ.get("DSH_LO_NODE") or LO_NODE_DEFAULT
    cli = cli_hint or os.environ.get("DSH_LO_CLI") or LO_CLI_DEFAULT
    if not os.path.isfile(cli):
        found = None
        for root in (os.path.dirname(os.path.dirname(os.path.dirname(__file__))),
                     os.getcwd()):
            cand = os.path.join(root, "node_modules", "@deepseek-ai", "libreoffice-kit",
                                "lib", "cli.js")
            if os.path.isfile(cand):
                found = cand
                break
        if not found:
            return None
        cli = found
    if not os.path.isfile(node):
        alt = shutil.which("node") or shutil.which("node.exe")
        if not alt:
            return None
        node = alt
    return (node, cli)


def _probe_soffice():
    cand = shutil.which("soffice") or shutil.which("soffice.exe") or shutil.which("libreoffice")
    if cand:
        return cand
    for path in SOFFICE_CANDIDATES:
        if os.path.isfile(path):
            return path
    return None


def _probe_word():
    try:
        import win32com.client  # noqa: F401
        import pythoncom        # noqa: F401
        return True
    except Exception:
        return False


def resolve_engine(requested, node_hint=None, cli_hint=None, warnings=None):
    """返回 (name, converter)。converter(docx, pdf) -> None。"""
    requested = (requested or "auto").lower()

    def lo_conv(node, cli):
        def _conv(docx_path, pdf_path, timeout=300):
            cmd = [node, cli, "convert", "--input", docx_path, "--output", pdf_path]
            proc = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                                  timeout=timeout)
            out = proc.stdout.decode("utf-8", "replace").strip()
            err = proc.stderr.decode("utf-8", "replace").strip()
            if proc.returncode != 0:
                raise EngineError("LibreOffice Kit 转换失败（rc=%d）：%s" % (proc.returncode, err or out))
            try:
                info = json.loads(out.splitlines()[-1]) if out else {}
            except Exception:
                info = {}
            if warnings is not None:
                miss = info.get("missingFonts") or []
                if miss:
                    warnings.append("LibreOffice 报告缺失字体：%s" % "、".join(miss))
        return _conv

    def soffice_conv(binary):
        def _conv(docx_path, pdf_path, timeout=300):
            outdir = os.path.dirname(os.path.abspath(pdf_path))
            cmd = [binary, "--headless", "--norestore", "--convert-to", "pdf",
                   "--outdir", outdir, docx_path]
            proc = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                                  timeout=timeout)
            made = os.path.join(outdir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")
            if not os.path.isfile(made):
                raise EngineError("soffice 未产出 PDF（rc=%d）：%s"
                                  % (proc.returncode, proc.stderr.decode("utf-8", "replace")[:400]))
            if os.path.abspath(made) != os.path.abspath(pdf_path):
                shutil.move(made, pdf_path)
        return _conv

    def word_conv(docx_path, pdf_path, timeout=300):
        import pythoncom
        from win32com.client import DispatchEx
        pythoncom.CoInitialize()
        app = None
        try:
            app = DispatchEx("Word.Application")
            app.Visible = False
            app.DisplayAlerts = 0
            d = app.Documents.Open(os.path.abspath(docx_path), ReadOnly=True,
                                   AddToRecentFiles=False)
            try:
                d.ExportAsFixedFormat(os.path.abspath(pdf_path), 17)   # wdExportFormatPDF
            finally:
                d.Close(0)
        finally:
            if app is not None:
                try:
                    app.Quit()
                except Exception:
                    pass
            pythoncom.CoUninitialize()

    if requested == "chrome":
        raise EngineError("chrome 引擎不支持 docx→pdf（无页脚页码域、无表格表头重复），"
                          "请改用 --engine libreoffice 或 word")

    lo = _probe_lo_kit(node_hint, cli_hint)
    if requested == "libreoffice":
        if not lo:
            raise EngineError("找不到 LibreOffice Kit：请设置 DSH_LO_NODE / DSH_LO_CLI，"
                              "或用 --lo-node / --lo-cli 指定 node 与 cli.js 路径")
        return "libreoffice", lo_conv(*lo)
    if requested == "soffice":
        exe = _probe_soffice()
        if not exe:
            raise EngineError("找不到 soffice：请安装 LibreOffice 或改用 --engine libreoffice")
        return "soffice", soffice_conv(exe)
    if requested == "word":
        if not _probe_word():
            raise EngineError("Word COM 不可用：需要安装 Microsoft Word 与 pywin32")
        return "word", word_conv

    # auto：libreoffice → soffice → word
    if lo:
        return "libreoffice", lo_conv(*lo)
    exe = _probe_soffice()
    if exe:
        return "soffice", soffice_conv(exe)
    if _probe_word():
        return "word", word_conv
    raise EngineError("没有可用的 PDF 转换引擎：请安装 LibreOffice（推荐，或让 DSH LibreOffice "
                      "Kit 可用）或 Microsoft Word + pywin32")


# =====================================================================================
# 7. PDF 校验
# =====================================================================================
_TOOL_CACHE = {}


def _tool_path(name):
    """按 PATH → TeX Live → poppler 常见目录的顺序找 pdfinfo/pdftoppm/pdftotext。"""
    if name in _TOOL_CACHE:
        return _TOOL_CACHE[name]
    hit = shutil.which(name)
    if not hit:
        for root in (r"D:\texlive", r"C:\Program Files\poppler", r"C:\poppler"):
            if hit or not os.path.isdir(root):
                continue
            for dirpath, dirnames, filenames in os.walk(root):
                if os.path.basename(dirpath).lower() in ("windows", "bin") and \
                        (name + ".exe") in filenames:
                    hit = os.path.join(dirpath, name + ".exe")
                    break
                if dirpath.count(os.sep) - root.count(os.sep) > 2:
                    dirnames[:] = []
    _TOOL_CACHE[name] = hit
    return hit


def pdf_pages(path):
    """页数：pdfinfo 优先，退回 /Type /Pages 的 /Count，再退回 /Type /Page 计数。"""
    info = _tool_path("pdfinfo")
    if info:
        try:
            proc = subprocess.run([info, path], stdout=subprocess.PIPE,
                                  stderr=subprocess.PIPE, timeout=60)
            m = re.search(rb"Pages:\s+(\d+)", proc.stdout)
            if m:
                return int(m.group(1)), "pdfinfo"
        except Exception:
            pass
    try:
        data = open(path, "rb").read()
    except OSError:
        return 0, "unreadable"
    counts = re.findall(rb"/Type\s*/Pages[^>]*?/Count\s+(\d+)", data)
    if counts:
        return max(int(c) for c in counts), "scan:/Count"
    return len(re.findall(rb"/Type\s*/Page[^s]", data)), "scan:/Type /Page"


def pdf_images(path):
    """/Subtype /Image 计数（近似值，失败不致命）。

    带 alpha 的插图会被拆成「图像 + smask」两个图像对象，减掉 SMask 引用
    才是插图张数。
    """
    try:
        data = open(path, "rb").read()
    except OSError:
        return 0
    total = len(re.findall(rb"/Subtype\s*/Image", data))
    masks = len(re.findall(rb"/SMask\s+\d+\s+\d+\s+R", data))
    return max(0, total - masks)


def validate_pdf(path):
    if not os.path.isfile(path):
        raise EngineError("转换后找不到 PDF：%s" % path)
    size = os.path.getsize(path)
    if size < 1024:
        raise EngineError("PDF 体积异常（%d 字节），转换可能失败" % size)
    with open(path, "rb") as fh:
        head = fh.read(8)
    if not head.startswith(b"%PDF-"):
        raise EngineError("PDF 签名校验失败，文件头为 %r" % head)
    pages, how = pdf_pages(path)
    if pages < 1:
        raise EngineError("PDF 页数异常（%d，统计方式 %s）" % (pages, how))
    return size, pages


# =====================================================================================
# 8. CLI
# =====================================================================================
def load_meta(path, warnings):
    meta = {}
    if not path:
        warnings.append("未提供 --meta：按无元数据处理（封面只有标题占位）")
        return meta
    if not os.path.isfile(path):
        warnings.append("meta 文件不存在：%s（按无元数据处理）" % path)
        return meta
    with open(path, encoding="utf-8") as fh:
        try:
            meta = json.load(fh)
        except json.JSONDecodeError as exc:
            raise EngineError("meta JSON 解析失败（%s）：%s" % (path, exc))
    if not isinstance(meta, dict):
        raise EngineError("meta 必须是 JSON 对象：%s" % path)
    meta["_meta_dir"] = os.path.dirname(os.path.abspath(path))
    return meta


def build_parser():
    ap = argparse.ArgumentParser(
        prog="build_report_pdf.py",
        description="Markdown 研报 → 中文 A4 DOCX → PDF（契约见 docs/report-pipeline.md）")
    ap.add_argument("--md", required=True, help="Markdown 正文路径")
    ap.add_argument("--out-dir", required=True, help="输出目录（生成 <name>.docx / <name>.pdf）")
    ap.add_argument("--meta", help="report-meta.json（封面与页脚元数据）")
    ap.add_argument("--engine", default="auto",
                    choices=["auto", "libreoffice", "soffice", "word", "chrome"],
                    help="转换引擎，默认 auto")
    ap.add_argument("--title", help="覆盖 meta.title")
    ap.add_argument("--name", default="report", help="输出文件基名，默认 report")
    ap.add_argument("--json", action="store_true", help="stdout 只输出一行 JSON")
    ap.add_argument("--no-docx", dest="no_docx", action="store_true",
                    help="转换后删除中间 docx")
    ap.add_argument("--keep-docx", dest="keep_docx", action="store_true",
                    help="保留中间 docx（默认行为，兼容契约 §6 的别名）")
    ap.add_argument("--max-image-width-cm", dest="img_max_w_cm", type=float,
                    default=IMG_MAX_W_CM, help="插图最大宽度（cm），默认 15.6")
    ap.add_argument("--lo-node", help="LibreOffice Kit 的 node 可执行文件")
    ap.add_argument("--lo-cli", help="LibreOffice Kit 的 cli.js")
    ap.add_argument("--timeout", type=float, default=300.0, help="单次转换超时（秒）")
    ap.add_argument("--h2-pagebreak", choices=["never", "always"], default="never",
                    help="## 是否强制段前分页：never（默认，连续排版，避免半空页）"
                         " / always（契约 §4 原始行为）")
    return ap


def _emit(payload, as_json, human_lines):
    if as_json:
        sys.stdout.write(json.dumps(payload, ensure_ascii=False) + "\n")
        return
    for line in human_lines:
        sys.stdout.write(line + "\n")


def main(argv=None):
    args = build_parser().parse_args(argv)
    warnings = []
    out_dir = os.path.abspath(args.out_dir)
    md_path = os.path.abspath(args.md)
    stage = "parse"

    payload = {"ok": False, "docx": "", "pdf": "", "pages": 0, "bytes": 0,
               "images": 0, "engine": "", "warnings": warnings}
    try:
        if not os.path.isfile(md_path):
            raise EngineError("找不到 Markdown 文件：%s" % md_path)
        os.makedirs(out_dir, exist_ok=True)

        with open(md_path, encoding="utf-8") as fh:
            raw = fh.read()
        raw, dropped = sanitize(raw)
        if dropped:
            warnings.append("已丢弃 %d 个中文字体无字形的表情符号" % dropped)

        meta = load_meta(args.meta, warnings)
        if args.title:
            meta["title"] = args.title

        blocks, parse_warnings = parse_markdown(raw)
        warnings.extend(parse_warnings)
        if not blocks:
            raise EngineError("Markdown 解析后没有任何内容：%s" % md_path)

        # ---- docx ----
        stage = "docx"
        doc, figures = build_document(meta, blocks, os.path.dirname(md_path),
                                      warnings, args.img_max_w_cm,
                                      h2_pagebreak=(args.h2_pagebreak == "always"))
        docx_path = os.path.join(out_dir, args.name + ".docx")
        doc.save(docx_path)
        if not os.path.isfile(docx_path) or os.path.getsize(docx_path) < 2048:
            raise EngineError("docx 写出异常：%s" % docx_path)

        # ---- convert ----
        stage = "convert"
        engine_name, convert = resolve_engine(args.engine, args.lo_node, args.lo_cli, warnings)
        pdf_path = os.path.join(out_dir, args.name + ".pdf")
        if os.path.isfile(pdf_path):
            os.remove(pdf_path)                       # 避免读到上一轮的旧产物
        t0 = time.time()
        convert(docx_path, pdf_path, timeout=args.timeout)
        took = time.time() - t0

        # ---- validate ----
        stage = "validate"
        size, pages = validate_pdf(pdf_path)
        images = pdf_images(pdf_path)
        if figures and images == 0:
            warnings.append("正文有 %d 张图但 PDF 未检出内嵌图片，请目检确认" % figures)

        if args.no_docx and not args.keep_docx:
            try:
                os.remove(docx_path)
                docx_path = ""
            except OSError:
                pass
        elif os.path.getsize(docx_path) < 2048:
            pass

        payload.update({"ok": True, "docx": docx_path, "pdf": pdf_path, "pages": pages,
                        "bytes": size, "images": images, "engine": engine_name})
        human = [
            "OK  %s" % pdf_path,
            "    engine=%s  pages=%d  bytes=%d  images=%d  figures=%d  用时=%.1fs"
            % (engine_name, pages, size, images, figures, took),
            "    docx=%s" % (docx_path or "（--no-docx 已删除）"),
            "    md=%s  blocks=%d  图片上限=%.1fcm" % (md_path, len(blocks), args.img_max_w_cm),
        ]
        for w in warnings:
            human.append("    warn: %s" % w)
        _emit(payload, args.json, human)
        return 0

    except EngineError as exc:
        payload.update({"ok": False, "stage": stage, "error": str(exc)})
        sys.stderr.write("[build_report_pdf] 阶段失败：%s → %s\n" % (stage, exc))
        if args.json:
            sys.stdout.write(json.dumps(payload, ensure_ascii=False) + "\n")
        return 1
    except subprocess.TimeoutExpired as exc:
        payload.update({"ok": False, "stage": stage, "error": "转换超时：%s" % exc})
        sys.stderr.write("[build_report_pdf] 阶段失败：%s → 转换超时（%s 秒）\n"
                         % (stage, args.timeout))
        if args.json:
            sys.stdout.write(json.dumps(payload, ensure_ascii=False) + "\n")
        return 1
    except Exception as exc:                          # noqa: BLE001 - CLI 边界要给出阶段
        payload.update({"ok": False, "stage": stage, "error": "%s: %s"
                        % (type(exc).__name__, exc)})
        sys.stderr.write("[build_report_pdf] 阶段失败：%s → %s: %s\n"
                         % (stage, type(exc).__name__, exc))
        if args.json:
            sys.stdout.write(json.dumps(payload, ensure_ascii=False) + "\n")
        return 1


if __name__ == "__main__":
    sys.exit(main())



